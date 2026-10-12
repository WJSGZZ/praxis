"""Independent small-table oracles and parsing boundaries, without case data."""
import copy
import hashlib
import json
import os
from pathlib import Path
import random
import subprocess
import sys

import pandas as pd
import pytest
from openpyxl import Workbook

from modeling.data_analysis import audit_join
from scripts.audit_data import audit


def _enumerated_join(left, right, keys, how):
    """Small independent Cartesian enumeration; used only as a test oracle."""
    pairs = [(i, j) for i, row_l in enumerate(left) for j, row_r in enumerate(right)
             if all(row_l[key] == row_r[key] for key in keys)]
    matched_l, matched_r = {i for i, _ in pairs}, {j for _, j in pairs}
    unmatched_l, unmatched_r = set(range(len(left))) - matched_l, set(range(len(right))) - matched_r
    output = pairs.copy()
    if how in {"left", "outer"}:
        output += [(i, None) for i in sorted(unmatched_l)]
    if how in {"right", "outer"}:
        output += [(None, j) for j in sorted(unmatched_r)]
    return output, matched_l, matched_r


@pytest.mark.parametrize("how", ["inner", "left", "right", "outer"])
def test_join_matches_independent_enumeration_and_preserves_input(how):
    left = [{"id": "a", "site": 1, "payload": 10}, {"id": "a", "site": 1, "payload": 20},
            {"id": "b", "site": 2, "payload": 30}, {"id": "c", "site": 2, "payload": 40}]
    right = [{"id": "a", "site": 1, "payload": 90}, {"id": "a", "site": 1, "payload": 100},
             {"id": "b", "site": 2, "payload": 110}, {"id": "d", "site": 2, "payload": 120}]
    before = copy.deepcopy((left, right))
    output, matched_l, matched_r = _enumerated_join(left, right, ["id", "site"], how)
    report = audit_join(left, right, ["id", "site"], how=how, relationship="many_to_many")
    assert report["output_rows"] == report["estimated_output_rows"] == len(output)
    assert report["indicator_counts"]["both"] == 5
    assert report["left"]["matched_input_rows"] == len(matched_l) == 3
    assert report["right"]["matched_input_rows"] == len(matched_r) == 3
    for side, position in [("left", 0), ("right", 1)]:
        included = [pair[position] for pair in output if pair[position] is not None]
        assert report[side]["retained_input_rows"] == len(set(included))
        assert report[side]["output_row_contributions"] == len(included)
        assert report[side]["duplicate_key_rows"] == 2
        assert report[side]["duplicate_key_excess_rows"] == 1
        assert report[side]["unmatched_input_rows"] == 1
    assert report["overlapping_payload_columns"] == ["payload"]
    assert (left, right) == before
    assert "records" not in report and "payload" not in json.dumps(report["left"])
    json.dumps(report, allow_nan=False)


@pytest.mark.parametrize("relationship,left,right", [
    ("many_to_one", [{"id": 1}], [{"id": 1}, {"id": 1}]),
    ("one_to_many", [{"id": 1}, {"id": 1}], [{"id": 1}]),
    ("one_to_one", [{"id": 1}], [{"id": 1}, {"id": 1}]),
])
def test_pandas_rejects_false_relationship(relationship, left, right):
    with pytest.raises(pd.errors.MergeError):
        audit_join(left, right, "id", relationship=relationship)


def test_expansion_bound_is_enforced_before_row_pair_materialization(monkeypatch):
    original_merge = pd.DataFrame.merge
    called_lengths = []
    def guarded_merge(frame, right, *args, **kwargs):
        called_lengths.append((len(frame), len(right)))
        if len(frame) > 1 or len(right) > 1:
            pytest.fail("An unbounded row-pair merge was attempted")
        return original_merge(frame, right, *args, **kwargs)
    monkeypatch.setattr(pd.DataFrame, "merge", guarded_merge)
    with pytest.raises(ValueError, match="1002001 rows.*no row-pair join"):
        audit_join([{"id": 1}] * 1001, [{"id": 1}] * 1001, "id", relationship="many_to_many")
    assert called_lengths == [(1, 1)]


def test_wide_key_join_has_a_cell_bound_before_row_pair_materialization(monkeypatch):
    keys = [f"key_{index}" for index in range(60)]
    row = {key: index for index, key in enumerate(keys)}
    original_merge = pd.DataFrame.merge
    calls = []
    def guarded_merge(frame, right, *args, **kwargs):
        calls.append((len(frame), len(right)))
        if len(frame) > 1 or len(right) > 1:
            pytest.fail("A wide unbounded row-pair join was attempted")
        return original_merge(frame, right, *args, **kwargs)
    monkeypatch.setattr(pd.DataFrame, "merge", guarded_merge)
    with pytest.raises(ValueError, match="63000000 key/row-ID cells.*no row-pair join"):
        audit_join([row] * 1000, [row] * 1000, keys, relationship="many_to_many", max_output_rows=1_000_000)
    assert calls == [(1, 1)]


def test_null_policy_and_original_effective_denominators():
    left = [{"id": None}, {"id": 1}, {"id": 2}]
    right = [{"id": None}, {"id": 1}]
    with pytest.raises(ValueError, match="Null join keys rejected"):
        audit_join(left, right, "id")
    dropped = audit_join(left, right, "id", null_policy="drop")
    assert dropped["output_rows"] == 2
    assert dropped["left"]["original_rows"] == 3 and dropped["left"]["effective_rows"] == 2
    assert dropped["left"]["dropped_null_key_rows"] == 1
    assert dropped["left"]["unmatched_fraction_effective"] == .5
    matched = audit_join(left, right, "id", null_policy="match")
    assert matched["indicator_counts"]["both"] == 2 and matched["output_rows"] == 3
    assert matched["left"]["dropped_null_key_rows"] == 0
    partial = audit_join([{"id": None, "site": "a"}, {"id": None, "site": "b"}],
                         [{"id": None, "site": "a"}], ["id", "site"], null_policy="match")
    assert partial["left"]["matched_input_rows"] == 1


def test_large_integer_null_and_reserved_column_names_remain_exact():
    report = audit_join([{"__praxis_left_row": 2**53+1}, {"__praxis_left_row": None}],
                        [{"__praxis_left_row": 2**53}, {"__praxis_left_row": None}],
                        "__praxis_left_row", null_policy="match")
    assert report["indicator_counts"]["both"] == report["left"]["matched_input_rows"] == 1
    assert report["left"]["unmatched_input_rows"] == 1


@pytest.mark.parametrize("left,right,options", [
    ([{"id": 1}], [{"id": "1"}], {}),
    ([{"id": 1}, {"id": 1.0}], [{"id": 1}], {}),
    ([{"id": True}], [{"id": 1}], {}),
    ([{"id": float("nan")}], [{"id": 1}], {}),
    ([{"id": float("inf")}], [{"id": 1}], {}),
    ([{"id": [1]}], [{"id": 1}], {}),
    ([{"id": 1, "value": {"bad": float("inf")}}], [{"id": 1}], {}),
    ([{"id": 1}, {"other": 1}], [{"id": 1}], {}),
    ([{"id": 1}], [{"id": 1}], {"max_output_rows": True}),
    ([{"id": 1}], [{"id": 1}], {"max_output_rows": 1_000_001}),
    ([{"id": 1}], [{"id": 1}], {"how": []}),
    ([{"id": 1}], [{"id": 1}], {"relationship": "m:1"}),
    ([{"id": 1}], [{"id": 1}], {"null_policy": "ignore"}),
])
def test_join_rejects_invalid_values_schema_and_policy(left, right, options):
    with pytest.raises(ValueError):
        audit_join(left, right, "id", **options)


def test_empty_join_reports_unknown_rates_and_bounds():
    report = audit_join([], [], "id")
    assert report["output_rows"] == 0 and report["left"]["expansion_factor"] is None
    report = audit_join([], [{"id": "a"}], "id", how="outer")
    assert report["output_rows"] == 1 and report["indicator_counts"]["right_only"] == 1


def test_composite_null_and_duplicate_joins_against_seeded_enumeration():
    for seed in range(12):
        random_source = random.Random(seed)
        left = [{"entity": random_source.choice([None, "a", "b", "c"]), "site": random_source.choice([None, 1, 2])} for _ in range(seed % 5)]
        right = [{"entity": random_source.choice([None, "a", "b", "d"]), "site": random_source.choice([None, 1, 2])} for _ in range(seed % 4)]
        for how in ["inner", "left", "right", "outer"]:
            output, matched_l, matched_r = _enumerated_join(left, right, ["entity", "site"], how)
            report = audit_join(left, right, ["entity", "site"], how=how,
                                relationship="many_to_many", null_policy="match")
            assert report["output_rows"] == len(output)
            assert report["left"]["matched_input_rows"] == len(matched_l)
            assert report["right"]["matched_input_rows"] == len(matched_r)
            output_keys = {tuple((left[i] if i is not None else right[j])[key] for key in ["entity", "site"]) for i, j in output}
            assert report["key_entities"]["output"] == len(output_keys)


@pytest.mark.parametrize("extension,separator", [("csv", ","), ("tsv", "\t")])
def test_duplicate_original_headers_are_disclosed_positionally_and_bytes_preserved(tmp_path, extension, separator):
    path = tmp_path / f"source.{extension}"
    content = separator.join(["id", "value", "value", "value.1", ""]) + "\n" + separator.join(["1", "2", "3", "4", "5"]) + "\n"
    original = content.encode()
    path.write_bytes(original)
    report = audit(path)
    assert report["header"]["original_names"] == ["id", "value", "value", "value.1", ""]
    assert report["header"]["parsed_names"] == ["id", "value", "value.2", "value.1", "Unnamed: 4"]
    assert report["header"]["duplicate_original_names"] == [{"name": "value", "positions": [1, 2]}]
    assert report["header"]["mapping"][2] == {"position": 2, "original": "value", "parsed": "value.2", "renamed": True}
    assert report["numeric_summary"]["value.2"]["mean"] == 3
    assert report["sha256"] == hashlib.sha256(original).hexdigest() and path.read_bytes() == original
    assert report["original_preserved"] and report["scan"]["total_rows"] == 1
    assert report["data_contract"]["metadata"]["population"]["status"] == "unknown"
    assert report["data_contract"]["column_metadata"][2]["unit"]["status"] == "unknown"


def test_quoted_blank_bom_and_multiline_original_headers(tmp_path):
    path = tmp_path / "quoted.csv"
    path.write_bytes(b'\xef\xbb\xbf"","multi\nline","a,b","a,b"\n1,2,3,4\n')
    report = audit(path)
    assert report["header"]["original_names"] == ["", "multi\nline", "a,b", "a,b"]
    assert report["header"]["duplicate_original_names"] == [{"name": "a,b", "positions": [2, 3]}]


def test_xlsx_headers_sheet_index_and_file_bytes_preserved(tmp_path):
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active.append(["id", "value", "value"])
    workbook.active.append([1, 2, 3])
    workbook.save(path)
    original = path.read_bytes()
    report = audit(path, sheet="0")
    assert report["header"]["original_names"] == ["id", "value", "value"]
    assert report["header"]["parsed_names"] == ["id", "value", "value.1"]
    assert report["parser"]["sheet"] == 0 and path.read_bytes() == original
    assert audit(path, sheet="Data")["rows"] == 1


def test_contract_failures_remain_diagnostic_and_prefix_scope_is_explicit(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text("id,mass,group\n1,-2,a\n1,,z\n2,8,b\n", encoding="utf-8")
    contract = {"observation_unit": "one measurement", "columns": {
        "mass": {"meaning": "Declared mass", "unit": "kg", "kind": "numeric", "nullable": False, "minimum": 0, "maximum": 5},
        "group": {"allowed_values": ["a", "b"]}}, "unique_keys": [["id"]]}
    before = path.read_bytes()
    report = audit(path, contract=contract, max_rows=2)
    assert report["rows"] == 2 and report["numeric_summary"]["mass"]["mean"] == -2
    assert report["data_contract"]["status"] == "issues_found"
    assert report["data_contract"]["unique_key_checks"][0]["duplicate_key_rows"] == 2
    assert {(issue.get("column"), issue["check"]) for issue in report["data_contract"]["issues"]} == {("mass", "not_null"), ("mass", "minimum"), ("group", "allowed_values"), (None, "unique_key")}
    assert report["scan"]["scope"] == "row_prefix" and report["scan"]["total_rows"] is None
    assert report["scan"]["rows_read"] == 3 and report["scan"]["rows_available_lower_bound"] == 3
    assert "not a random" in report["scan"]["sampling"] and path.read_bytes() == before
    assert audit(path, max_rows=3)["scan"]["scope"] == "full_table"


def test_prefix_lookahead_does_not_change_inference_and_index_inference_is_disclosed(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text("value\n1\n2\nnot-a-number\n")
    prefix = audit(path, max_rows=2)
    assert prefix["numeric_summary"]["value"]["mean"] == 1.5
    assert prefix["dtypes"]["value"] == "int64" and prefix["scan"]["truncated"]
    path.write_text("a,b\n1,2,3\n4,5,6\n")
    implicit = audit(path)
    assert implicit["parser"]["implicit_index_inference"]
    assert implicit["rows"] == 2 and implicit["columns"] == 2


def test_explicit_encoding_delimiter_missing_markers_and_old_positional_interface(tmp_path):
    path = tmp_path / "source.csv"
    path.write_bytes("名称;value\n甲;NA\n乙;-999\n丙;3\n".encode("gb18030"))
    legacy = audit(path, 0, "gb18030", ";")
    assert legacy["missing"]["value"] == 1 and legacy["parser"]["keep_default_na"] is True
    literal = audit(path, encoding="gb18030", sep=";", keep_default_na=False, missing_values=["-999"])
    assert literal["missing"]["value"] == 1 and literal["parser"]["missing_values_declared"] == ["-999"]
    assert literal["dtypes"]["value"] == "object"
    assert literal["parser"]["keep_default_na"] is False


def test_nonfinite_observations_are_disclosed_without_changing_them_to_missing(tmp_path):
    path = tmp_path / "source.csv"
    original = b"value\ninf\n-inf\nNaN\n1\n"
    path.write_bytes(original)
    report = audit(path, contract={"columns": {"value": {"kind": "numeric", "finite": True}}})
    assert report["rows"] == 4 and report["missing"]["value"] == 1
    assert report["column_profiles"][0]["nonfinite_nonmissing"] == 2
    assert report["numeric_summary"]["value"]["count"] == 3
    states = report["numeric_summary_status"]["value"]["null_serialized_statistics"]
    assert states["max"] == "positive_infinity" and states["min"] == "negative_infinity"
    finite = report["numeric_finite_summary"]["value"]
    assert finite["included_finite"] == 1 and finite["excluded_nonfinite_nonmissing"] == 2
    assert finite["excluded_missing"] == 1 and finite["summary"]["mean"] == 1
    assert report["data_contract"]["issues"][0]["check"] == "finite"
    assert report["data_contract"]["issues"][0]["count"] == 2
    assert path.read_bytes() == original
    prefix = audit(path, max_rows=1)
    assert prefix["scan"]["scope"] == "row_prefix"
    assert prefix["column_profiles"][0]["nonfinite_nonmissing"] == 1
    assert prefix["numeric_finite_summary"]["value"]["status"] == "no_finite_observations"
    assert prefix["data_contract"]["status"] == "not_declared"


def test_numeric_summary_overflow_is_separate_from_nonfinite_input(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text("value\n1e308\n1e308\n")
    report = audit(path)
    assert report["column_profiles"][0]["nonfinite_nonmissing"] == 0
    assert report["numeric_summary"]["value"]["mean"] is None
    assert report["numeric_summary_status"]["value"]["null_serialized_statistics"]["mean"] == "positive_infinity"
    assert report["numeric_finite_summary"] == {}


@pytest.mark.parametrize("options", [{"max_rows": 0}, {"max_rows": True}, {"sep": "||"},
                                      {"keep_default_na": "no"}, {"missing_values": {"x": 1}},
                                      {"contract": {"columns": {"x": {"coerce": True}}}},
                                      {"contract": {"columns": {"x": {"kind": []}}}},
                                      {"contract": {"columns": {"x": {"allowed_values": None}}}},
                                      {"contract": {"columns": {"x": {"minimum": float("nan")}}}}])
def test_audit_rejects_invalid_configuration(tmp_path, options):
    path = tmp_path / "source.csv"
    path.write_text("x\n1\n")
    with pytest.raises(ValueError):
        audit(path, **options)


def test_cli_audits_real_file_and_rejects_overwriting_source(tmp_path):
    path = tmp_path / "source.tsv"
    path.write_text("id\tx\tx\n1\t2\t3\n")
    command = [sys.executable, "-m", "scripts.audit_data", str(path)]
    completed = subprocess.run(command, capture_output=True, text=True, check=True)
    assert json.loads(completed.stdout)["header"]["automatic_renaming"]
    before = path.read_bytes()
    refused = subprocess.run(command + ["--output", str(path)], capture_output=True, text=True)
    assert refused.returncode != 0 and "must not overwrite input" in refused.stderr
    assert path.read_bytes() == before


def test_cli_rejects_existing_hardlink_output_without_changing_original(tmp_path):
    source, alias = tmp_path / "source.csv", tmp_path / "alias.csv"
    original = b"id,value\n1,2\n2,3\n"
    source.write_bytes(original)
    try:
        os.link(source, alias)
    except (OSError, NotImplementedError) as error:
        pytest.skip(f"Hardlinks unavailable on this test filesystem: {error}")
    assert source.samefile(alias) and source.resolve() != alias.resolve()
    invocation = subprocess.run([sys.executable, "-m", "scripts.audit_data", str(source), "--output", str(alias)],
                                capture_output=True, text=True)
    assert invocation.returncode != 0 and "must not overwrite input" in invocation.stderr
    assert source.read_bytes() == alias.read_bytes() == original
    target = tmp_path / "new-output" / "audit.json"
    success = subprocess.run([sys.executable, "-m", "scripts.audit_data", str(source), "--output", str(target)],
                             capture_output=True, text=True, check=True)
    assert json.loads(success.stdout)["sha256"] == hashlib.sha256(original).hexdigest()
    assert json.loads(target.read_text())["rows"] == 2 and source.read_bytes() == original


def test_cli_still_rejects_symlink_output(tmp_path):
    source, alias = tmp_path / "source.csv", tmp_path / "alias.csv"
    original = b"id,value\n1,2\n"
    source.write_bytes(original)
    try:
        alias.symlink_to(source)
    except (OSError, NotImplementedError) as error:
        pytest.skip(f"Symlinks unavailable on this test filesystem: {error}")
    invocation = subprocess.run([sys.executable, "-m", "scripts.audit_data", str(source), "--output", str(alias)],
                                capture_output=True, text=True)
    assert invocation.returncode != 0 and "must not overwrite input" in invocation.stderr
    assert source.read_bytes() == original


def test_numeric_xlsx_header_contract_resolves_one_json_display_name(tmp_path):
    path = tmp_path / "numeric-headers.xlsx"
    workbook = Workbook()
    workbook.active.append([101, 2026])
    workbook.active.append([1, 5])
    workbook.active.append([1, 8])
    workbook.save(path)
    original = path.read_bytes()
    contract = {"columns": {"101": {"meaning": "Entity ID", "unit": "identifier", "kind": "integer"},
                            "2026": {"meaning": "Observed measure", "unit": "kg", "maximum": 6}},
                "unique_keys": [["101"]]}
    report = audit(path, contract=contract)
    assert report["header"]["original_names"] == [101, 2026]
    assert report["header"]["parsed_names"] == [101, 2026]
    references = report["data_contract"]["column_matches"]
    assert references[0] == {"display_name": "101", "status": "matched", "matches": [{"position": 0, "parsed_name": 101, "parsed_name_type": "int"}]}
    limit = next(check for check in report["data_contract"]["checks"] if check["check"] == "maximum")
    assert limit["status"] == "failed" and limit["count"] == 1
    assert limit["column_reference"]["matches"][0]["position"] == 1
    assert report["data_contract"]["unique_key_checks"][0]["duplicate_key_rows"] == 2
    assert report["data_contract"]["column_metadata"][1]["unit"] == {"status": "declared", "value": "kg"}
    assert report["legacy_key_collisions"] == [] and path.read_bytes() == original
    prefix = audit(path, contract=contract, max_rows=1)
    assert prefix["scan"]["scope"] == "row_prefix"
    assert prefix["data_contract"]["unique_key_checks"][0]["status"] == "passed"


def test_mixed_typed_xlsx_names_are_ambiguous_and_all_positional_summaries_survive(tmp_path):
    path = tmp_path / "mixed-headers.xlsx"
    workbook = Workbook()
    workbook.active.append([101, "101", 2026])
    workbook.active.append([1, 10, 5])
    workbook.active.append([3, None, 8])
    workbook.save(path)
    original = path.read_bytes()
    report = audit(path, contract={"columns": {"101": {"meaning": "Unresolved", "unit": "kg", "maximum": 0}},
                                  "unique_keys": [["101"]]})
    assert report["header"]["original_names"] == [101, "101", 2026]
    assert report["header"]["parsed_names"] == [101, "101", 2026]
    assert report["header"]["original_name_types"] == ["int", "str", "int"]
    reference = report["data_contract"]["column_matches"][0]
    assert reference["status"] == "ambiguous"
    assert [match["parsed_name_type"] for match in reference["matches"]] == ["int", "str"]
    assert [match["position"] for match in reference["matches"]] == [0, 1]
    assert report["data_contract"]["checks"] == []
    assert report["data_contract"]["issues"][0]["status"] == "not_checked"
    assert report["data_contract"]["unique_key_checks"][0]["status"] == "not_checked"
    assert [row["unit"]["status"] for row in report["data_contract"]["column_metadata"][:2]] == ["ambiguous", "ambiguous"]
    assert report["dtypes"]["101"] is None and report["missing"]["101"] is None
    assert report["numeric_summary"]["101"] is None and report["numeric_summary_status"]["101"] is None
    assert report["legacy_key_collisions"][0]["status"] == "ambiguous"
    profiles = report["column_profiles"]
    assert profiles[0]["missing"] == 0 and profiles[1]["missing"] == 1
    assert profiles[0]["numeric_summary"]["mean"] == 2
    assert profiles[1]["numeric_summary"]["mean"] == 10
    assert profiles[2]["numeric_summary"]["mean"] == 6.5
    assert path.read_bytes() == original


def test_direct_audit_script_remains_usable_outside_source_cwd(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text("id,x,x\n1,2,3\n")
    script = Path(__file__).resolve().parents[1] / "scripts" / "audit_data.py"
    run = subprocess.run([sys.executable, str(script), str(path)], cwd=tmp_path,
                         capture_output=True, text=True, check=True)
    report = json.loads(run.stdout)
    assert report["header"]["duplicate_original_names"] == [{"name": "x", "positions": [1, 2]}]
    assert report["numeric_summary"]["x.1"]["mean"] == 3


def test_store_demo_has_exact_answer_complete_family_and_preserved_inputs(tmp_path):
    from examples.data_analysis_demo import run_demo
    output = tmp_path / "store-demo"
    report = run_demo(output)
    regression = json.loads((output / "cluster-regression.json").read_text())
    assert [row["estimate"] for row in regression["coefficients"]] == pytest.approx([50, 7, 2], abs=1e-11)
    assert regression["covariance"]["clusters"] == 12 and regression["covariance"]["df_inference"] == 11
    assert regression["covariance"]["matrix"][1][1] == pytest.approx(94 / 27, abs=1e-11)
    family = json.loads((output / "comparison-family.json").read_text())
    assert [row["hypothesis_id"] for row in family["hypotheses"]] == ["campaign_association", "weekly_trend"]
    assert family["family_size"] == 2 and family["method"] == "holm"
    observed = pd.read_csv(output / "observations.csv")
    assert len(observed) == 48 and observed["store_id"].nunique() == 12
    assert not observed.duplicated(["store_id", "week"]).any()
    assert observed.groupby("store_id").size().tolist() == [4] * 12
    rejected = json.loads((output / "bad-join-rejection.json").read_text())
    assert rejected["status"] == "rejected" and rejected["exception"] == "MergeError"
    oracle = json.loads((output / "independent-verification.json").read_text())
    assert oracle["status"] == "passed" and oracle["original_inputs_preserved"]
    assert oracle["cluster_covariance_max_absolute_error"] < 1e-11
    for name, digest in report["original_input_sha256"].items():
        assert hashlib.sha256((output / name).read_bytes()).hexdigest() == digest
    before = (output / "report.json").read_bytes()
    with pytest.raises(FileExistsError):
        run_demo(output)
    assert (output / "report.json").read_bytes() == before


def test_store_demo_direct_cli_reproduces_identical_local_inputs(tmp_path):
    from examples.data_analysis_demo import run_demo
    first = run_demo(tmp_path / "first")
    second = tmp_path / "second"
    invocation = subprocess.run([sys.executable, "examples/data_analysis_demo.py", "--output", str(second)],
                                capture_output=True, text=True, check=True)
    assert json.loads(invocation.stdout)["status"] == "verified_synthetic_demo"
    reproduced = json.loads((second / "report.json").read_text())
    assert reproduced["original_input_sha256"] == first["original_input_sha256"]
